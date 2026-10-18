#!/usr/bin/env python3
r"""Run with the packaged interpreter on EACH target OS/architecture.

  <payload>/runtime/bin/python3 verify_company_runtime.py <payload>
  <payload>\runtime\python.exe verify_company_runtime.py <payload>

Produces an optional target verification record for offline packaging only after actual tool tests.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib
import importlib.metadata
import json
import platform
import runpy
import subprocess
import sys
import tempfile
from pathlib import Path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    args = parser.parse_args()
    root = args.root.resolve()
    runtime = root / "runtime"
    if not Path(sys.executable).resolve().is_relative_to(runtime):
        raise SystemExit("Run this check using the payload's Python interpreter, not system Python")
    sys.path.insert(0, str(root / "packages"))
    launcher = runpy.run_path(str(root / "launch.py"))
    launcher["_bootstrap"]()
    for module in (
        "octop.company_runtime",
        "octop_harness",
        "openpyxl",
        "docx",
        "pptx",
        "pypdf",
        "PIL",
        "pandas",
    ):
        importlib.import_module(module)
    harness_version = importlib.metadata.version("octop-harness")
    if harness_version != "1.0.3":
        raise SystemExit(f"Expected locked octop-harness 1.0.3, got {harness_version}")
    system = "windows" if sys.platform == "win32" else sys.platform
    arch = {"x86_64": "amd64", "amd64": "amd64", "aarch64": "arm64", "arm64": "arm64"}.get(
        platform.machine().lower()
    )
    if system not in ("windows", "darwin") or not arch:
        raise SystemExit("Enterprise packages support Windows/macOS amd64/arm64")
    suffix = ".exe" if system == "windows" else ""
    ffmpeg, ffprobe = (root / "tools" / (name + suffix) for name in ("ffmpeg", "ffprobe"))
    for tool in (ffmpeg, ffprobe):
        subprocess.run([str(tool), "-version"], check=True, capture_output=True, timeout=20)
    with tempfile.TemporaryDirectory(prefix="octop-company-smoke-") as temp:
        folder = Path(temp)
        import pandas as pd
        from docx import Document
        from openpyxl import Workbook, load_workbook
        from PIL import Image
        from pptx import Presentation
        from pypdf import PdfReader, PdfWriter

        book = Workbook()
        book.active.append([2, 3])
        book.save(folder / "sample.xlsx")
        assert sum(load_workbook(folder / "sample.xlsx").active.values.__next__()) == 5
        assert pd.read_excel(folder / "sample.xlsx", header=None).sum().sum() == 5
        doc = Document()
        doc.add_paragraph("Octop")
        doc.save(folder / "sample.docx")
        assert Document(folder / "sample.docx").paragraphs[0].text == "Octop"
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(folder / "sample.pptx")
        assert len(Presentation(folder / "sample.pptx").slides) == 1
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.write(folder / "sample.pdf")
        assert len(PdfReader(folder / "sample.pdf").pages) == 1
        Image.new("RGB", (16, 16), "blue").save(folder / "sample.png")
        assert Image.open(folder / "sample.png").size == (16, 16)
        video = folder / "sample.mp4"
        subprocess.run(
            [
                str(ffmpeg),
                "-nostdin",
                "-y",
                "-f",
                "lavfi",
                "-i",
                "color=blue:s=64x64:d=0.2",
                "-c:v",
                "mpeg4",
                str(video),
            ],
            check=True,
            capture_output=True,
            timeout=30,
        )
        subprocess.run(
            [str(ffprobe), "-v", "error", "-show_streams", str(video)],
            check=True,
            capture_output=True,
            timeout=20,
        )
        subprocess.run(
            [
                sys.executable,
                "-c",
                "from pathlib import Path; Path('terminal-ok.txt').write_text('ok')",
            ],
            cwd=folder,
            check=True,
            timeout=20,
        )
        assert (folder / "terminal-ok.txt").read_text() == "ok"
    hashes = {}
    for item in sorted(root.rglob("*")):
        if not item.is_file() or "__pycache__" in item.parts:
            continue
        relative = item.relative_to(root).as_posix()
        if relative in ("company-runtime-build.json", "company-runtime-manifest.json"):
            continue
        with item.open("rb") as stream:
            hashes[relative] = hashlib.file_digest(stream, "sha256").hexdigest()
    from octop.infra.company_runtime import PROTOCOL_VERSION as company_protocol
    from octop.infra.devices.protocol import PROTOCOL_VERSION as device_protocol

    record = {
        "octop_version": importlib.metadata.version("octop"),
        "device_protocol_version": device_protocol,
        "company_protocol_version": company_protocol,
        "file_hashes": hashes,
        "protocol_version": 1,
        "platform": system,
        "arch": arch,
        "python_version": platform.python_version(),
        "harness_version": harness_version,
        "smoke_test_passed": True,
    }
    (root / "company-runtime-build.json").write_text(json.dumps(record, indent=2) + "\n")
    print(json.dumps(record))


if __name__ == "__main__":
    main()
