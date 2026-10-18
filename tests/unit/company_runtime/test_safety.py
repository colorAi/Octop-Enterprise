import asyncio
import hashlib
import json
import os
import shutil
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest

from octop.infra.company_runtime.client import CompanyServiceError
from octop.infra.company_runtime.local import ProjectBackend

from .test_runtime import acquire, install_agent


@pytest.mark.asyncio
async def test_restart_rebind_requires_installation_secret_and_never_steals(company):
    server, client, credentials, owner = company
    install_agent(server)
    await acquire(client)
    client.http.headers["X-Octop-Device-Credential"] = credentials[1]
    recovery = {"installation_token": "i" * 32, "recovery_token": "x" * 32}
    with pytest.raises(CompanyServiceError, match="previous_device_still_active"):
        await client.request("tasks/task1/recover", recovery)
    server.services.device_repo.revoke("device0", owner, 1)
    with pytest.raises(CompanyServiceError, match="recovery_not_owned"):
        await client.request("tasks/task1/recover", {**recovery, "installation_token": "z" * 32})
    state = await client.request("tasks/task1/recover", recovery)
    assert state["status"] == "unknown"
    with pytest.raises(CompanyServiceError, match="thread_already_running"):
        await acquire(client, "task2")
    await client.request("tasks/task1/status", {"status": "failed"})
    assert (await acquire(client, "task2"))["status"] == "running"
    await client.close()


@pytest.mark.asyncio
async def test_legacy_and_server_local_execution_never_mix(company):
    server, client, _, _ = company
    install_agent(server)
    server.services.thread_repo.touch_last_active("thread1")
    with pytest.raises(CompanyServiceError, match="legacy_thread_read_only"):
        await acquire(client)
    await acquire(client, "task2", "thread2")
    with pytest.raises(ValueError, match="company_local_thread_read_only_on_server"):
        server.services.thread_repo.claim_server_execution("thread2")
    server.services.thread_repo.claim_server_execution("thread3")
    with pytest.raises(CompanyServiceError, match="legacy_thread_read_only"):
        await acquire(client, "task3", "thread3")
    await client.close()


@pytest.mark.asyncio
async def test_real_company_memory_is_idempotent_and_cross_session(company, tmp_path):
    from octop_memory import Memory

    server, client, _, _ = company
    install_agent(server)
    memory = Memory(
        namespace="runtime_test",
        backend="sqlite",
        backend_config={"db_path": str(tmp_path / "company_memory.sqlite")},
    )
    server.app_runtime.agent_registry.get_agent("expert").memory = memory
    await acquire(client)
    body = {"operation_id": "remember1", "assertion": "The employee prefers teal charts."}
    first = await client.request("tasks/task1/memory/remember", body)
    assert (await client.request("tasks/task1/memory/remember", body)) == first
    await acquire(client, "task2", "thread2")
    found = await client.request("tasks/task2/memory/search", {"query": "teal", "max_results": 5})
    assert "teal" in json.dumps(found)
    await client.close()
    memory._backend.close()


@pytest.mark.asyncio
async def test_image_generation_receipt_and_download_scope(company, monkeypatch):
    server, client, _, _ = company
    install_agent(server)
    data = b"generated-test-image"
    agent = server.app_runtime.agent_registry.get_agent("expert")
    agent.config.media_generation = SimpleNamespace(has_api_key=lambda: True, image_enabled=True)
    agent.workspace.adownload_bytes = AsyncMock(return_value=data)
    generate = AsyncMock(
        return_value=SimpleNamespace(
            artifacts=[SimpleNamespace(path="generated/cover.png", media_type="image/png")]
        )
    )
    monkeypatch.setattr(
        "octop.infra.company_runtime.media.create_media_manager",
        lambda *args: SimpleNamespace(generate_image=generate),
    )
    await acquire(client)
    request = {"operation_id": "image1", "prompt": "a cover"}
    result = await client.request("tasks/task1/images", request)
    assert (await client.request("tasks/task1/images", request)) == result
    assert generate.await_count == 1
    item = result["artifacts"][0]
    response = await client.http.get(f"tasks/task1/images/image1/{item['file_id']}")
    assert response.content == data
    assert response.headers["x-content-sha256"] == hashlib.sha256(data).hexdigest()
    assert (await client.http.get("tasks/task1/images/image1/not-owned")).status_code == 404
    await client.close()


@pytest.mark.asyncio
async def test_image_download_failure_is_reported_without_regenerating(company, monkeypatch):
    from octop_harness.media.errors import MediaGenerationError

    server, client, _, _ = company
    install_agent(server)
    agent = server.app_runtime.agent_registry.get_agent("expert")
    agent.config.media_generation = SimpleNamespace(has_api_key=lambda: True, image_enabled=True)
    generate = AsyncMock(
        side_effect=MediaGenerationError(
            "Generated media download timed out.",
            code="provider_timeout",
            remediation="retry_download",
        )
    )
    monkeypatch.setattr(
        "octop.infra.company_runtime.media.create_media_manager",
        lambda *args: SimpleNamespace(generate_image=generate),
    )
    await acquire(client)
    for _ in range(2):
        with pytest.raises(CompanyServiceError, match="company_image_download_failed") as error:
            await client.request(
                "tasks/task1/images", {"operation_id": "image1", "prompt": "cover"}
            )
        assert error.value.status == 502
    assert generate.await_count == 1
    assert server.services.company_runtime_repo.operation("task1", "image1")["status"] == "failed"
    await client.close()


@pytest.mark.asyncio
async def test_image_request_has_its_own_long_wait_budget():
    import httpx

    from octop.infra.company_runtime.client import CompanyClient

    timeouts = []

    def handle(request):
        timeouts.append(request.extensions["timeout"])
        return httpx.Response(200, json={})

    client = CompanyClient("http://company.test", "test", transport=httpx.MockTransport(handle))
    await client.request("tasks/task1/images", {"operation_id": "test", "prompt": "cover"})
    await client.request("capabilities")
    assert timeouts[0] == {"connect": 10, "read": 900, "write": 900, "pool": 900}
    assert timeouts[1]["read"] == 120
    await client.close()


@pytest.mark.asyncio
async def test_background_pipe_respects_timeout(tmp_path):
    if os.name == "nt":
        pytest.skip("POSIX background-process syntax")
    backend = ProjectBackend(tmp_path, tmp_path / "logs", lambda event: None, tool_path=os.defpath)
    result = await asyncio.wait_for(backend.aexecute("sleep 30 &", timeout=1), 3)
    assert result.exit_code == 124
    assert not backend.processes


@pytest.mark.asyncio
async def test_local_excel_docx_and_ffmpeg_with_unicode_paths(tmp_path):
    import openpyxl
    from docx import Document

    project = tmp_path / "中文 project"
    project.mkdir()
    book = openpyxl.Workbook()
    book.active.append(["value"])
    book.active.append([3])
    book.active.append([7])
    book.save(project / "source.xlsx")
    script = project / "calculate.py"
    script.write_text(
        "import openpyxl\nfrom docx import Document\nw=openpyxl.load_workbook('source.xlsx',data_only=True)\ntotal=sum(row[0] for row in list(w.active.values)[1:])\nd=Document();d.add_paragraph(str(total));d.save('report.docx')\n",
        encoding="utf-8",
    )
    backend = ProjectBackend(
        project,
        tmp_path / "state",
        lambda event: None,
        tool_path=os.pathsep.join([str(Path(sys.executable).parent), os.defpath]),
    )
    result = await backend.aexecute(f'"{sys.executable}" "{script}"')
    assert result.exit_code == 0, result.output
    assert Document(project / "report.docx").paragraphs[0].text == "10"
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        pytest.skip("FFmpeg unavailable on test host")
    result = await backend.aexecute(
        f'"{ffmpeg}" -v error -f lavfi -i color=c=blue:s=160x90:d=1 -c:v libx264 "input.mp4"'
    )
    assert result.exit_code == 0, result.output
    before = hashlib.sha256((project / "input.mp4").read_bytes()).hexdigest()
    result = await backend.aexecute(
        f'"{ffmpeg}" -v error -i "input.mp4" -vf scale=80:46 -c:v libx264 "output.mp4"'
    )
    assert result.exit_code == 0, result.output
    assert (project / "output.mp4").stat().st_size > 1000
    assert hashlib.sha256((project / "input.mp4").read_bytes()).hexdigest() == before


@pytest.mark.asyncio
async def test_concurrent_project_isolation_and_native_absolute_paths(company, tmp_path):
    from octop.infra.company_runtime.runtime import EnterpriseRuntime

    server, client, _, _ = company
    install_agent(server)
    runtime = EnterpriseRuntime(lambda event: None)
    runtime.client, runtime.workspace, runtime.state = client, tmp_path / "work", tmp_path / "state"
    runtime.workspace.mkdir()
    runtime.authorized, runtime.installation_token = True, "i" * 32
    runtime.tools = {"ffmpeg": "test", "openpyxl": "test", "python-docx": "test", "Pillow": "test"}
    for number in (1, 2):
        await runtime.start(
            {
                "task_id": f"task{number}",
                "thread_id": f"thread{number}",
                "agent_id": "expert",
                "text": "create a report",
            }
        )
    await asyncio.gather(*list(runtime.work.values()))
    assert all(t["status"] == "completed" for t in runtime.tasks.values())
    assert len(list(runtime.workspace.rglob("report.txt"))) == 2
    project = runtime.workspace / "thread1"
    backend = ProjectBackend(project, tmp_path / "logs", lambda event: None, tool_path=os.defpath)
    assert backend._resolve_path(str(project / "脚本 space.py")) == project / "脚本 space.py"
    with pytest.raises(ValueError):
        backend._resolve_path("../escape")
    await runtime.close()


@pytest.mark.asyncio
async def test_other_employee_cannot_use_expert_or_checkpoints(company):
    server, client, credentials, owner = company
    install_agent(server)
    other = server.services.user_repo.create(username="other", role="admin")
    # Even an admin device must not act as another employee's private expert.
    with server.services.db.transaction() as conn:
        conn.execute(
            "UPDATE desktop_devices SET owner_user_id=? WHERE device_id='device1'", (other,)
        )
    await acquire(client)
    client.http.headers["X-Octop-Device-Credential"] = credentials[1]
    with pytest.raises(CompanyServiceError, match="expert_not_owned"):
        await acquire(client, "other_task", "thread2")
    response = await client.http.post("tasks/task1/checkpoint/get", json={})
    assert response.status_code == 403
    from dataclasses import replace

    server.services = replace(
        server.services, config=replace(server.services.config, company_runtime_enabled=False)
    )
    assert (await client.http.get("capabilities")).status_code == 404
    await client.close()


@pytest.mark.asyncio
async def test_skill_assets_are_pinned_and_permission_scoped(company, tmp_path):
    from deepagents.backends import FilesystemBackend
    from octop_harness.backends.workspace import BackendWorkspace
    from octop_harness.skills.catalog import list_skill_summaries

    server, client, _, _ = company
    install_agent(server)
    source = tmp_path / "expert-source"
    skill = source / "skills" / "calculate"
    skill.mkdir(parents=True)
    (skill / "SKILL.md").write_text(
        "---\nname: calculate\ndescription: Calculate values\n---\nUse run.py."
    )
    script = skill / "run.py"
    script.write_text("print(42)")
    workspace = BackendWorkspace(FilesystemBackend(root_dir=source, virtual_mode=True), source)
    agent = server.app_runtime.agent_registry.get_agent("expert")
    agent.workspace = workspace
    agent.config.skills_dir = None
    agent.list_skill_summaries = lambda: list_skill_summaries(workspace)
    task = await acquire(client)
    assets = task["expert"]["skills"]
    assert any(item["path"] == "skills/calculate/SKILL.md" for item in assets)
    code = next(item for item in assets if item["path"] == "skills/calculate/run.py")
    script.write_text("print(100)")
    response = await client.http.get(f"tasks/task1/skills/{code['sha256']}")
    assert response.content == b"print(42)"
    assert (await client.http.get("tasks/task1/skills/" + "a" * 64)).status_code == 403
    next_task = await acquire(client, "task2", "thread2")
    assert next_task["expert"]["version"] != task["expert"]["version"]
    await client.close()


@pytest.mark.asyncio
async def test_result_index_survives_a_new_repository_instance(company):
    from octop.infra.db.repos.company_runtime import CompanyRuntimeRepo

    server, client, _, _ = company
    install_agent(server)
    await acquire(client)
    files = [{"path": "/original-computer/report.docx", "size": 42, "sha256": "f" * 64}]
    await client.request(
        "tasks/task1/events",
        {
            "event_id": "user1",
            "seq": 1,
            "payload": {
                "kind": "messages",
                "messages": [
                    {"type": "human", "data": {"id": "task1_user", "content": "create report"}}
                ],
            },
        },
    )
    await client.request(
        "tasks/task1/events",
        {
            "event_id": "result1",
            "seq": 2,
            "payload": {"kind": "artifacts", "location": "desktop", "files": files},
        },
    )
    await client.request("tasks/task1/status", {"status": "completed"})
    repo = CompanyRuntimeRepo(server.services.db)
    assert repo.artifacts("thread1") == files
    assert repo.artifacts("thread2") == []
    await client.close()


@pytest.mark.asyncio
async def test_expert_deletion_cleans_runtime_children(company):
    server, client, _, _ = company
    install_agent(server)
    await acquire(client)
    await client.request(
        "tasks/task1/events", {"event_id": "event1", "seq": 1, "payload": {"kind": "note"}}
    )
    server.services.agent_repo.delete("expert")
    with server.services.db.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM company_runtime_threads").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM company_runtime_tasks").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM company_runtime_events").fetchone()[0] == 0
    await client.close()
